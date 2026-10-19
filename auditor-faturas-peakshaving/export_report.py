"""
export_report.py
----------------
Geração de relatório técnico-financeiro em formato Excel (.xlsx) profissional
com openpyxl, contendo:
1. Resumo Executivo & Auditoria de Conformidade (Cativo ACR vs. Livre ACL)
2. Memória de Cálculo Completa Passo a Passo (Normas ANEEL)
3. Tabela Exaustiva de Todas as Tarifas Utilizadas (Com e Sem Tributos)
4. Auditoria Fiscal Detalhada dos Tributos (ICMS, PIS, COFINS, CIP e Gross-up)
"""

from __future__ import annotations

import io
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from invoice_parser import InvoiceData
from tariff_client import ConcessionariaTarifas
from tariff_engine import ParametrosGerador, ResultadoSimulacao


def gerar_relatorio_excel(
    resultado: ResultadoSimulacao,
    fatura: InvoiceData,
    gerador: ParametrosGerador,
    tarifas_homologadas: ConcessionariaTarifas,
) -> bytes:
    """
    Constrói um arquivo .xlsx em memória com 4 abas estruturadas e formatadas profissionalmente:
    1. Resumo Executivo & Auditoria
    2. Memória de Cálculo Detalhada
    3. Tarifas Utilizadas
    4. Tributos Detalhados
    """
    wb = openpyxl.Workbook()

    # Estilos Visuais Padronizados
    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=11, bold=True, color="1E293B")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=10, color="1E293B")
    font_kpi_val = Font(name="Calibri", size=12, bold=True, color="059669")
    font_audit_ok = Font(name="Calibri", size=10, bold=True, color="059669")
    font_audit_alert = Font(name="Calibri", size=10, bold=True, color="DC2626")

    fill_header_navy = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_header_green = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    fill_header_blue = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    fill_subtle_gray = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_accent_green = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    fill_accent_blue = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    fill_accent_alert = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

    thin_border_side = Side(style="thin", color="CBD5E1")
    border_all = Border(
        left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    num_fmt_currency = "R$ #,##0.00"
    num_fmt_decimal4 = "0.0000"
    num_fmt_decimal5 = "0.00000"
    num_fmt_integer = "#,##0"
    num_fmt_percent = "0.0%"

    # =========================================================================
    # ABA 1: RESUMO EXECUTIVO & AUDITORIA DA FATURA
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Resumo & Auditoria"
    ws1.views.sheetView[0].showGridLines = True

    # Banner Superior
    ws1.merge_cells("A1:G1")
    c_title = ws1["A1"]
    c_title.value = "AUDITORIA TARIFÁRIA E VIABILIDADE DE PEAK SHAVING - GRUPO A4"
    c_title.font = font_title
    c_title.fill = fill_header_navy
    c_title.alignment = align_center
    ws1.row_dimensions[1].height = 34

    ws1.merge_cells("A2:G2")
    c_sub = ws1["A2"]
    c_sub.value = (
        f"Distribuidora: {resultado.distribuidora} | "
        f"Ambiente: {resultado.ambiente_contratacao} (ACR/ACL) | "
        f"Norma Regulamentar: {resultado.resolucao_homologatoria} | "
        f"Vigência: {resultado.data_vigencia or 'Atual'}"
    )
    c_sub.font = Font(name="Calibri", size=9.5, italic=True, color="E2E8F0")
    c_sub.fill = fill_header_navy
    c_sub.alignment = align_center
    ws1.row_dimensions[2].height = 20

    # 1. BLOCO DE AUDITORIA DE CONFORMIDADE DA FATURA
    ws1.cell(row=4, column=1, value="1. AUDITORIA DE CONFORMIDADE DA FATURA ATUAL (CONFRONTO ANEEL)").font = font_section

    audit_headers = ["Parâmetro Auditado", "Valor Fatura Cobrado", "Valor ANEEL Reconstituído", "Diferença / Desvio (R$)", "Desvio (%)", "Diagnóstico de Conformidade", "Ambiente"]
    for col_idx, h in enumerate(audit_headers, 1):
        cell = ws1.cell(row=5, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = align_center
        cell.border = border_all
    ws1.row_dimensions[5].height = 24

    ws1.row_dimensions[6].height = 22
    ws1.cell(row=6, column=1, value="Fatura Concessionária (Rede + Encargos)").font = font_bold
    ws1.cell(row=6, column=2, value=resultado.valor_fatura_cobrado_real).number_format = num_fmt_currency
    ws1.cell(row=6, column=3, value=resultado.valor_fatura_reconstituida_reais).number_format = num_fmt_currency
    ws1.cell(row=6, column=4, value=resultado.desvio_auditoria_reais).number_format = num_fmt_currency
    ws1.cell(row=6, column=5, value=resultado.desvio_auditoria_percentual / 100.0).number_format = num_fmt_percent
    c_diag = ws1.cell(row=6, column=6, value="CONFORME (< 2.5% variação)" if resultado.fatura_auditada_conforme else "DIVERGENTE")
    c_diag.font = font_audit_ok if resultado.fatura_auditada_conforme else font_audit_alert
    c_diag.alignment = align_center
    ws1.cell(row=6, column=7, value=f"{resultado.ambiente_contratacao} (ACR/ACL)").alignment = align_center

    for c_i in range(1, 8):
        ws1.cell(row=6, column=c_i).border = border_all
        if c_i in [2, 3, 4, 5]:
            ws1.cell(row=6, column=c_i).alignment = align_right

    # Parecer textual da auditoria
    ws1.merge_cells("A7:G8")
    c_parecer = ws1["A7"]
    c_parecer.value = f"Parecer da Auditoria: {resultado.diagnostico_auditoria}"
    c_parecer.font = Font(name="Calibri", size=9.5, italic=True, color="1E293B")
    c_parecer.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c_parecer.fill = fill_accent_green if resultado.fatura_auditada_conforme else fill_accent_alert
    ws1.row_dimensions[7].height = 18
    ws1.row_dimensions[8].height = 18

    # 2. BLOCO DE INDICADORES PRINCIPAIS DE DESEMPENHO (KPIS)
    ws1.cell(row=10, column=1, value="2. COMPARATIVO OPERACIONAL DOS CENÁRIOS (KPIS)").font = font_section

    kpi_headers = ["Indicador Operacional", "Valor Atual (Cenário 0)", "Verde + Gerador (Cen. 1)", "Azul + Gerador (Cen. 2)", "Economia Líquida", "% Redução"]
    for col_idx, h in enumerate(kpi_headers, 1):
        cell = ws1.cell(row=11, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = align_center
        cell.border = border_all
    ws1.row_dimensions[11].height = 24

    kpi_rows = [
        ("Despesa Operacional Total Mensal (R$)", resultado.cenario_0_atual.custo_total_operacional_reais, resultado.cenario_1_verde.custo_total_operacional_reais, resultado.cenario_2_azul.custo_total_operacional_reais, resultado.economia_mensal_liquida_reais, resultado.economia_mensal_percentual / 100.0),
        ("Fatura Concessionária / Distribuidora (R$)", resultado.cenario_0_atual.custo_fatura_concessionaria_reais, resultado.cenario_1_verde.custo_fatura_concessionaria_reais, resultado.cenario_2_azul.custo_fatura_concessionaria_reais, resultado.cenario_0_atual.custo_fatura_concessionaria_reais - min(resultado.cenario_1_verde.custo_fatura_concessionaria_reais, resultado.cenario_2_azul.custo_fatura_concessionaria_reais), (resultado.cenario_0_atual.custo_fatura_concessionaria_reais - min(resultado.cenario_1_verde.custo_fatura_concessionaria_reais, resultado.cenario_2_azul.custo_fatura_concessionaria_reais)) / resultado.cenario_0_atual.custo_fatura_concessionaria_reais if resultado.cenario_0_atual.custo_fatura_concessionaria_reais > 0 else 0),
        ("Contrato Bilateral Livre ACL (R$)", resultado.cenario_0_atual.custo_energia_livre_reais, resultado.cenario_1_verde.custo_energia_livre_reais, resultado.cenario_2_azul.custo_energia_livre_reais, resultado.cenario_0_atual.custo_energia_livre_reais - resultado.cenario_1_verde.custo_energia_livre_reais, 0.0),
        ("Gasto Operacional do Skid Gerador (R$)", 0.0, resultado.cenario_1_verde.gasto_gerador_total_reais, resultado.cenario_2_azul.gasto_gerador_total_reais, -min(resultado.cenario_1_verde.gasto_gerador_total_reais, resultado.cenario_2_azul.gasto_gerador_total_reais), 0.0),
        ("Consumo Ativo Ponta na Rede (kWh)", resultado.cenario_0_atual.consumo_ponta_rede_kwh, resultado.cenario_1_verde.consumo_ponta_rede_kwh, resultado.cenario_2_azul.consumo_ponta_rede_kwh, resultado.cenario_0_atual.consumo_ponta_rede_kwh - resultado.cenario_1_verde.consumo_ponta_rede_kwh, 1.0 if resultado.cenario_1_verde.consumo_ponta_rede_kwh == 0 else (resultado.cenario_0_atual.consumo_ponta_rede_kwh - resultado.cenario_1_verde.consumo_ponta_rede_kwh) / resultado.cenario_0_atual.consumo_ponta_rede_kwh),
        ("Demanda Contratada Fora Ponta (kW)", resultado.cenario_0_atual.demanda_contratada_fora_ponta_kw, resultado.cenario_1_verde.demanda_contratada_fora_ponta_kw, resultado.cenario_2_azul.demanda_contratada_fora_ponta_kw, 0.0, 0.0),
        ("Demanda Contratada Ponta (kW)", resultado.cenario_0_atual.demanda_contratada_ponta_kw or 0.0, 0.0, resultado.cenario_2_azul.demanda_contratada_ponta_kw or 0.0, (resultado.cenario_0_atual.demanda_contratada_ponta_kw or 0.0) - (resultado.cenario_2_azul.demanda_contratada_ponta_kw or 0.0), 0.0),
    ]

    for r_idx, row_data in enumerate(kpi_rows, 12):
        ws1.row_dimensions[r_idx].height = 20
        for c_idx, val in enumerate(row_data, 1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_all
            if c_idx == 1:
                cell.alignment = align_left
                cell.font = font_bold
            elif c_idx in [2, 3, 4, 5]:
                cell.alignment = align_right
                if "R$" in row_data[0]:
                    cell.number_format = num_fmt_currency
                elif "kWh" in row_data[0] or "kW" in row_data[0]:
                    cell.number_format = num_fmt_integer
                if c_idx == 5:
                    cell.font = font_kpi_val
            elif c_idx == 6:
                cell.alignment = align_right
                cell.number_format = num_fmt_percent
                cell.font = font_bold

    # 3. CARD DE RECOMENDAÇÃO AUTOMÁTICA
    ws1.cell(row=20, column=1, value="3. DECISÃO ESTRATÉGICA E RECOMENDAÇÃO AUTOMÁTICA").font = font_section

    ws1.merge_cells("A21:G21")
    c_rec_title = ws1["A21"]
    c_rec_title.value = f"ENQUADRAMENTO RECOMENDADO: TARIFA HORÁRIA {resultado.modalidade_recomendada} COM PEAK SHAVING ({resultado.ambiente_contratacao})"
    c_rec_title.font = font_header
    c_rec_title.fill = fill_header_green if resultado.modalidade_recomendada == "VERDE" else fill_header_blue
    c_rec_title.alignment = align_center
    ws1.row_dimensions[21].height = 26

    rec_details = [
        ("Economia Mensal Líquida:", f"R$ {resultado.economia_mensal_liquida_reais:,.2f} ({resultado.economia_mensal_percentual:.1f}% de redução na despesa)"),
        ("Economia Anualizada Projetada:", f"R$ {resultado.economia_anual_projetada_reais:,.2f} ao ano"),
        ("Demanda Contratada Sugerida (FP):", f"{resultado.recomendacao_demanda_fp_kw:.1f} kW"),
        ("Demanda Contratada Sugerida (P):", f"{resultado.recomendacao_demanda_ponta_kw:.1f} kW" if resultado.recomendacao_demanda_ponta_kw else "Não aplicável (Tarifa Única na Verde)"),
        ("Custo Marginal do Gerador:", f"R$ {gerador.custo_gerador_unitario:.3f}/kWh (Diesel: R$ {gerador.preco_diesel_litro:.2f}/L | O&M: R$ {gerador.custo_manutencao_kwh:.2f}/kWh)"),
        ("Custo Unitário de Ponta Evitado:", f"R$ {resultado.comparativo_custo_ponta['tarifa_rede_verde_ponta_com_impostos']:.3f}/kWh (com tributos)"),
        ("Margem Unitária de Economia:", f"R$ {resultado.comparativo_custo_ponta['margem_unit_diesel_vs_rede_verde']:.3f} por kWh cortado no horário de ponta"),
    ]

    for idx, (label, val_str) in enumerate(rec_details, 22):
        ws1.row_dimensions[idx].height = 20
        c_lbl = ws1.cell(row=idx, column=1, value=label)
        c_lbl.font = font_bold
        c_lbl.alignment = align_left
        c_lbl.border = border_all
        c_lbl.fill = fill_subtle_gray

        ws1.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=7)
        c_v = ws1.cell(row=idx, column=2, value=val_str)
        c_v.font = font_regular
        c_v.alignment = align_left
        c_v.border = border_all

    # Justificativa Integral
    ws1.merge_cells("A30:G32")
    c_just = ws1["A30"]
    c_just.value = f"Justificativa Técnica: {resultado.justificativa_tecnica}"
    c_just.font = Font(name="Calibri", size=9.5, italic=True, color="334155")
    c_just.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c_just.fill = fill_accent_green if resultado.modalidade_recomendada == "VERDE" else fill_accent_blue

    # =========================================================================
    # ABA 2: MEMÓRIA DE CÁLCULO PASSO A PASSO
    # =========================================================================
    ws2 = wb.create_sheet(title="Memória de Cálculo")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:F1")
    c2_title = ws2["A1"]
    c2_title.value = f"MEMÓRIA DE CÁLCULO PASSO A PASSO - {resultado.ambiente_contratacao} (NORMAS ANEEL)"
    c2_title.font = font_title
    c2_title.fill = fill_header_navy
    c2_title.alignment = align_center
    ws2.row_dimensions[1].height = 30

    calc_headers = ["Item / Componente de Custo", "Fórmula / Base ANEEL", "Cenário 0 (Atual)", "Cenário 1 (Verde + Ger.)", "Cenário 2 (Azul + Ger.)", "Unidade"]
    for c_i, h in enumerate(calc_headers, 1):
        cell = ws2.cell(row=3, column=c_i, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = align_center
        cell.border = border_all
    ws2.row_dimensions[3].height = 22

    memoria_itens = [
        # DEMANDAS
        ("1. DEMANDA CONTRATADA E FATURADA", "", "", "", "", ""),
        ("Demanda Fora Ponta / Única", "Medição / Contrato", resultado.cenario_0_atual.demanda_contratada_fora_ponta_kw, resultado.cenario_1_verde.demanda_contratada_fora_ponta_kw, resultado.cenario_2_azul.demanda_contratada_fora_ponta_kw, "kW"),
        ("Demanda Ponta", "Medição / Otimizada", resultado.cenario_0_atual.demanda_contratada_ponta_kw or 0.0, 0.0, resultado.cenario_2_azul.demanda_contratada_ponta_kw or 0.0, "kW"),
        ("Tarifa Demanda Fora Ponta (Homologada)", "REH ANEEL Líquida", resultado.cenario_0_atual.tarifa_demanda_fp_homologada, resultado.cenario_1_verde.tarifa_demanda_fp_homologada, resultado.cenario_2_azul.tarifa_demanda_fp_homologada, "R$/kW"),
        ("Tarifa Demanda Ponta (Homologada)", "REH ANEEL Líquida", resultado.cenario_0_atual.tarifa_demanda_p_homologada or 0.0, 0.0, resultado.cenario_2_azul.tarifa_demanda_p_homologada or 0.0, "R$/kW"),
        ("Custo Total de Demanda (c/ Tributos)", "Demanda * Tarifa * Gross-up", resultado.cenario_0_atual.custo_demanda_total_reais, resultado.cenario_1_verde.custo_demanda_total_reais, resultado.cenario_2_azul.custo_demanda_total_reais, "R$"),
        # CONSUMOS
        ("2. CONSUMO ATIVO DE ENERGIA (TUSD & TE)", "", "", "", "", ""),
        ("Consumo Fora Ponta (Rede)", "Medição Ativa FP", resultado.cenario_0_atual.consumo_fora_ponta_kwh, resultado.cenario_1_verde.consumo_fora_ponta_kwh, resultado.cenario_2_azul.consumo_fora_ponta_kwh, "kWh"),
        ("Tarifa TUSD Fora Ponta (Homologada)", "REH ANEEL Líquida", resultado.cenario_0_atual.tarifa_tusd_fp_homologada, resultado.cenario_1_verde.tarifa_tusd_fp_homologada, resultado.cenario_2_azul.tarifa_tusd_fp_homologada, "R$/kWh"),
        ("Tarifa TE Fora Ponta (Concessionária)", "Zero no ACL / Homologada no ACR", resultado.cenario_0_atual.tarifa_te_fp_homologada, resultado.cenario_1_verde.tarifa_te_fp_homologada, resultado.cenario_2_azul.tarifa_te_fp_homologada, "R$/kWh"),
        ("Custo Consumo Fora Ponta (Rede)", "Consumo_FP * (TUSD_FP + TE_FP)", resultado.cenario_0_atual.custo_consumo_fp_total_reais, resultado.cenario_1_verde.custo_consumo_fp_total_reais, resultado.cenario_2_azul.custo_consumo_fp_total_reais, "R$"),
        ("Consumo Ponta Total Requerido", "Medição Ativa Ponta", resultado.cenario_0_atual.consumo_ponta_total_kwh, resultado.cenario_1_verde.consumo_ponta_total_kwh, resultado.cenario_2_azul.consumo_ponta_total_kwh, "kWh"),
        ("Consumo Ponta Suprido pelo Gerador", "min(Consumo_P, Potência*Dias*3h)", 0.0, resultado.cenario_1_verde.consumo_ponta_gerador_kwh, resultado.cenario_2_azul.consumo_ponta_gerador_kwh, "kWh"),
        ("Consumo Ponta Residual da Rede", "Consumo_P - Consumo_Gerador", resultado.cenario_0_atual.consumo_ponta_rede_kwh, resultado.cenario_1_verde.consumo_ponta_rede_kwh, resultado.cenario_2_azul.consumo_ponta_rede_kwh, "kWh"),
        ("Tarifa TUSD Ponta (Homologada)", "REH ANEEL Líquida", resultado.cenario_0_atual.tarifa_tusd_p_homologada, resultado.cenario_1_verde.tarifa_tusd_p_homologada, resultado.cenario_2_azul.tarifa_tusd_p_homologada, "R$/kWh"),
        ("Tarifa TE Ponta (Concessionária)", "Zero no ACL / Homologada no ACR", resultado.cenario_0_atual.tarifa_te_p_homologada, resultado.cenario_1_verde.tarifa_te_p_homologada, resultado.cenario_2_azul.tarifa_te_p_homologada, "R$/kWh"),
        ("Custo Consumo Ponta da Rede", "Consumo_P_Rede * (TUSD_P + TE_P)", resultado.cenario_0_atual.custo_consumo_ponta_rede_reais, resultado.cenario_1_verde.custo_consumo_ponta_rede_reais, resultado.cenario_2_azul.custo_consumo_ponta_rede_reais, "R$"),
        # MERCADO LIVRE ACL
        ("3. ENERGIA NO MERCADO LIVRE (ACL)", "", "", "", "", ""),
        ("Preço do Contrato Bilateral Livre", "PPA Negociado", resultado.cenario_0_atual.preco_energia_livre_kwh, resultado.cenario_1_verde.preco_energia_livre_kwh, resultado.cenario_2_azul.preco_energia_livre_kwh, "R$/kWh"),
        ("Custo Mensal do Contrato Livre", "Consumo_Rede * Preço_ACL", resultado.cenario_0_atual.custo_energia_livre_reais, resultado.cenario_1_verde.custo_energia_livre_reais, resultado.cenario_2_azul.custo_energia_livre_reais, "R$"),
        # GERADOR
        ("4. OPERAÇÃO DO GERADOR TÉRMICO (PEAK SHAVING)", "", "", "", "", ""),
        ("Preço do Óleo Diesel", "Parâmetro Operacional", "-", gerador.preco_diesel_litro, gerador.preco_diesel_litro, "R$/litro"),
        ("Consumo Específico", "Parâmetro Operacional", "-", gerador.consumo_especifico_l_kwh, gerador.consumo_especifico_l_kwh, "L/kWh"),
        ("Custo de Manutenção / O&M", "Parâmetro Operacional", "-", gerador.custo_manutencao_kwh, gerador.custo_manutencao_kwh, "R$/kWh"),
        ("Custo Marginal Unitário do Gerador", "(Preço_Diesel * Cons_Esp) + O&M", 0.0, gerador.custo_gerador_unitario, gerador.custo_gerador_unitario, "R$/kWh"),
        ("Gasto Mensal com Óleo Diesel", "Consumo_Ger * Cons_Esp * Preço", 0.0, resultado.cenario_1_verde.gasto_gerador_diesel_reais, resultado.cenario_2_azul.gasto_gerador_diesel_reais, "R$"),
        ("Gasto Mensal com O&M", "Consumo_Ger * Custo_Manut", 0.0, resultado.cenario_1_verde.gasto_gerador_manutencao_reais, resultado.cenario_2_azul.gasto_gerador_manutencao_reais, "R$"),
        ("Gasto Total Operacional do Gerador", "Gasto_Diesel + Gasto_O&M", 0.0, resultado.cenario_1_verde.gasto_gerador_total_reais, resultado.cenario_2_azul.gasto_gerador_total_reais, "R$"),
        # TOTAIS
        ("5. ENCARGOS ADICIONAIS E TOTAIS", "", "", "", "", ""),
        ("Bandeira Tarifária / Adicionais", "Isento no ACL / Rateio no ACR", resultado.cenario_0_atual.custo_bandeiras_reais, resultado.cenario_1_verde.custo_bandeiras_reais, resultado.cenario_2_azul.custo_bandeiras_reais, "R$"),
        ("Contribuição Iluminação Pública (CIP)", "Cobrança Municipal", resultado.cenario_0_atual.valor_cip_reais, resultado.cenario_1_verde.valor_cip_reais, resultado.cenario_2_azul.valor_cip_reais, "R$"),
        ("Custo Fatura Distribuidora (Rede)", "Demanda + TUSD + TE + CIP + Trib.", resultado.cenario_0_atual.custo_fatura_concessionaria_reais, resultado.cenario_1_verde.custo_fatura_concessionaria_reais, resultado.cenario_2_azul.custo_fatura_concessionaria_reais, "R$"),
        ("DESPESA OPERACIONAL LÍQUIDA TOTAL", "Fatura Distrib. + Contrato ACL + Gerador", resultado.cenario_0_atual.custo_total_operacional_reais, resultado.cenario_1_verde.custo_total_operacional_reais, resultado.cenario_2_azul.custo_total_operacional_reais, "R$"),
    ]

    for row_num, item in enumerate(memoria_itens, 4):
        ws2.row_dimensions[row_num].height = 19
        is_section_header = item[0].startswith(("1.", "2.", "3.", "4.", "5."))
        for c_i, val in enumerate(item, 1):
            cell = ws2.cell(row=row_num, column=c_i, value=val)
            cell.border = border_all
            if is_section_header:
                cell.font = font_bold
                cell.fill = fill_subtle_gray
            elif item[0] == "DESPESA OPERACIONAL LÍQUIDA TOTAL":
                cell.font = font_title
                cell.fill = fill_header_navy
            else:
                cell.font = font_regular

            if c_i == 1:
                cell.alignment = align_left
            elif c_i == 2 or c_i == 6:
                cell.alignment = align_center
            else:
                cell.alignment = align_right
                if isinstance(val, (int, float)):
                    if item[5] == "R$":
                        cell.number_format = num_fmt_currency
                    elif item[5] in ["R$/kWh", "R$/kW", "L/kWh", "R$/litro"]:
                        cell.number_format = num_fmt_decimal4
                    elif item[5] in ["kW", "kWh"]:
                        cell.number_format = num_fmt_integer

    # =========================================================================
    # ABA 3: TODAS AS TARIFAS UTILIZADAS (COM E SEM TRIBUTOS)
    # =========================================================================
    ws3 = wb.create_sheet(title="Tarifas Utilizadas")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:G1")
    c3_title = ws3["A1"]
    c3_title.value = f"TABELA EXAUSTIVA DE TARIFAS UTILIZADAS - {resultado.distribuidora} ({resultado.ambiente_contratacao})"
    c3_title.font = font_title
    c3_title.fill = fill_header_navy
    c3_title.alignment = align_center
    ws3.row_dimensions[1].height = 30

    t_headers = [
        "Componente Tarifário",
        "Unidade",
        "Modalidade",
        "Ambiente",
        "Tarifa ANEEL (Sem Tributos)",
        "Tarifa Efetiva (Com Tributos)",
        "Local de Faturamento / Incidência",
    ]
    for c_i, h in enumerate(t_headers, 1):
        cell = ws3.cell(row=3, column=c_i, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = align_center
        cell.border = border_all
    ws3.row_dimensions[3].height = 22

    for r_i, r_data in enumerate(resultado.tabela_tarifas_utilizadas, 4):
        ws3.row_dimensions[r_i].height = 20
        c1 = ws3.cell(row=r_i, column=1, value=r_data["Componente"])
        c2 = ws3.cell(row=r_i, column=2, value=r_data["Unidade"])
        c3 = ws3.cell(row=r_i, column=3, value=r_data["Modalidade"])
        c4 = ws3.cell(row=r_i, column=4, value=r_data["Ambiente"])
        c5 = ws3.cell(row=r_i, column=5, value=r_data["Tarifa_Homologada_Sem_Tributos"])
        c6 = ws3.cell(row=r_i, column=6, value=r_data["Tarifa_Efetiva_Com_Tributos"])
        c7 = ws3.cell(row=r_i, column=7, value=r_data["Onde_Faturado"])

        c1.alignment = align_left
        c2.alignment = align_center
        c3.alignment = align_center
        c4.alignment = align_center
        c5.alignment = align_right
        c6.alignment = align_right
        c7.alignment = align_left

        c5.number_format = num_fmt_decimal5
        c6.number_format = num_fmt_decimal5

        for col_idx in range(1, 8):
            cell = ws3.cell(row=r_i, column=col_idx)
            cell.border = border_all
            cell.font = font_bold if "Gerador" in r_data["Componente"] or "Demanda" in r_data["Componente"] else font_regular

    # =========================================================================
    # ABA 4: TRIBUTOS DETALHADOS (AUDITORIA FISCAL)
    # =========================================================================
    ws4 = wb.create_sheet(title="Tributos Detalhados")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:E1")
    c4_title = ws4["A1"]
    c4_title.value = f"DETALHAMENTO FISCAL DE TRIBUTOS E GROSS-UP - {resultado.distribuidora}"
    c4_title.font = font_title
    c4_title.fill = fill_header_navy
    c4_title.alignment = align_center
    ws4.row_dimensions[1].height = 30

    trib_headers = [
        "Tributo / Encargo",
        "Base de Cálculo (R$)",
        "Alíquota Nominal (%)",
        "Metodologia Legal de Cálculo",
        "Valor do Tributo Apurado (R$)",
    ]
    for c_i, h in enumerate(trib_headers, 1):
        cell = ws4.cell(row=3, column=c_i, value=h)
        cell.font = font_header
        cell.fill = fill_header_navy
        cell.alignment = align_center
        cell.border = border_all
    ws4.row_dimensions[3].height = 22

    for r_i, r_data in enumerate(resultado.tabela_tributos_utilizados, 4):
        ws4.row_dimensions[r_i].height = 20
        is_total = "TOTAL" in r_data["Tributo"]
        c1 = ws4.cell(row=r_i, column=1, value=r_data["Tributo"])
        c2 = ws4.cell(row=r_i, column=2, value=r_data["Base_Calculo_Reais"])
        c3 = ws4.cell(row=r_i, column=3, value=r_data["Aliquota_Nominal_Percent"] / 100.0 if r_data["Aliquota_Nominal_Percent"] > 0 else 0.0)
        c4 = ws4.cell(row=r_i, column=4, value=r_data["Metodologia"])
        c5 = ws4.cell(row=r_i, column=5, value=r_data["Valor_Apurado_Reais"])

        c1.alignment = align_left
        c2.alignment = align_right
        c3.alignment = align_right
        c4.alignment = align_left
        c5.alignment = align_right

        c2.number_format = num_fmt_currency
        c3.number_format = num_fmt_percent
        c5.number_format = num_fmt_currency

        for col_idx in range(1, 6):
            cell = ws4.cell(row=r_i, column=col_idx)
            cell.border = border_all
            if is_total:
                cell.font = font_bold
                cell.fill = fill_subtle_gray
            else:
                cell.font = font_regular

    # Ajuste automático de largura de colunas em todas as abas
    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if cell.number_format == num_fmt_currency:
                    max_len = max(max_len, len(val) + 6)
                else:
                    max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
